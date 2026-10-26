"""
Agente IA de Aldair · Ingeniería de Confiabilidad
=================================================

Funciona de dos formas con el MISMO archivo:

1) EN TU PC (gratis, privado, con Ollama)
       python -m pip install streamlit pypdf python-docx openpyxl
       ollama pull qwen2.5:14b
       ollama pull nomic-embed-text
       python -m streamlit run pagina.py

2) EN LA NUBE (Streamlit Community Cloud + Gemini de Google)
   Se activa sola cuando existe el secreto GEMINI_API_KEY. Secretos en Streamlit Cloud:
       GEMINI_API_KEY = "tu_clave_de_google_ai_studio"
       CLAVE_ADMIN    = "contraseña para ver Ajustes y Conocimiento"   (modo compañeros)
       MODELO_GEMINI  = "gemini-3.8-flash"                             (opcional)

Archivos que se crean junto a la página:
    conocimiento/                 documentos de la base de conocimiento
    conocimiento_indice.json      índice de búsqueda
    confiabot_ajustes.json        ajustes de la página
    confiabot_logo_empresa.png    logo de la empresa (opcional)
"""
from __future__ import annotations

import json
import os
import sys
import urllib.error
import urllib.request

import math

SYSTEM_PROMPT = """Eres un agente de inteligencia artificial personal que actúa como ingeniero de confiabilidad senior (CMRP, CRE) con más de 20 años
de experiencia en industrias de petróleo y gas, minería, energía, puertos y manufactura.

Dominas:
- Mantenimiento Centrado en Confiabilidad (RCM, SAE JA1011/JA1012).
- FMEA / FMECA (IEC 60812, AIAG-VDA), Análisis Causa Raíz (RCA, 5 porqués, Ishikawa, árbol lógico).
- Análisis estadístico de vida: Weibull, exponencial, lognormal; curva de la bañera; análisis de Crow-AMSAA.
- Indicadores: MTBF, MTTR, MTTF, disponibilidad, OEE, confiabilidad R(t), tasa de falla.
- Taxonomía y recolección de datos de falla según ISO 14224; terminología de EN 13306 e ISO 55000.
- Análisis de criticidad, RBD (diagramas de bloques de confiabilidad), costo del ciclo de vida (LCC),
  optimización de intervalos de mantenimiento preventivo y predictivo (CBM, curva P-F).

Reglas de trabajo:
1. Responde siempre en español, de forma clara y práctica, como un consultor en planta.
2. Cuando haya datos numéricos, USA LAS HERRAMIENTAS para calcular; nunca inventes resultados.
3. Explica qué significa cada resultado para la decisión de mantenimiento
   (p. ej. beta < 1 → mortalidad infantil: revisar montaje, arranque y calidad de repuestos;
   beta ≈ 1 → fallas aleatorias: el preventivo por tiempo no ayuda, preferir CBM;
   beta > 1 → desgaste: el reemplazo preventivo puede ser rentable).
4. Si faltan datos, indica cuáles necesitas y propone supuestos razonables, señalándolos como supuestos.
5. Advierte cuando la muestra sea pequeña (< 6 fallas) o los resultados tengan baja confianza.
6. Cita la norma o el método aplicable cuando sea pertinente.
"""


# ---------------------------------------------------------------------------
# Herramientas de cálculo (Python puro, sin dependencias externas)
# ---------------------------------------------------------------------------
def calcular_indicadores(tiempos_entre_fallas: list[float], tiempos_reparacion: list[float],
                         horas_periodo: float | None = None) -> dict:
    """MTBF, MTTR, disponibilidad inherente y tasa de falla."""
    if not tiempos_entre_fallas or not tiempos_reparacion:
        return {"error": "Se requieren tiempos entre fallas y tiempos de reparación."}
    n_fallas = len(tiempos_entre_fallas)
    mtbf = sum(tiempos_entre_fallas) / n_fallas
    mttr = sum(tiempos_reparacion) / len(tiempos_reparacion)
    disponibilidad = mtbf / (mtbf + mttr)
    res = {
        "numero_fallas": n_fallas,
        "MTBF_h": round(mtbf, 2),
        "MTTR_h": round(mttr, 2),
        "disponibilidad_inherente_%": round(disponibilidad * 100, 3),
        "tasa_falla_por_hora": round(1 / mtbf, 6),
        "tasa_falla_por_anio": round(8760 / mtbf, 3),
    }
    if horas_periodo:
        res["horas_fuera_servicio"] = round(sum(tiempos_reparacion), 2)
        res["disponibilidad_operacional_%"] = round((horas_periodo - sum(tiempos_reparacion)) / horas_periodo * 100, 3)
    return res


def ajuste_weibull(tiempos_falla: list[float], tiempos_suspension: list[float] | None = None) -> dict:
    """Ajuste Weibull de 2 parámetros por regresión de rangos de mediana (Bernard),
    con rangos ajustados de Johnson cuando hay suspensiones (datos censurados)."""
    tiempos_suspension = tiempos_suspension or []
    datos = sorted([(t, True) for t in tiempos_falla] + [(t, False) for t in tiempos_suspension])
    n = len(datos)
    if len(tiempos_falla) < 2:
        return {"error": "Se necesitan al menos 2 tiempos de falla."}
    if any(t <= 0 for t, _ in datos):
        return {"error": "Todos los tiempos deben ser mayores que cero."}

    # Rangos ajustados de Johnson
    puntos, rango_prev = [], 0.0
    for i, (t, es_falla) in enumerate(datos):
        if es_falla:
            rango_inverso = n - i
            incremento = (n + 1 - rango_prev) / (1 + rango_inverso)
            rango = rango_prev + incremento
            rango_prev = rango
            mr = (rango - 0.3) / (n + 0.4)            # aproximación de Bernard
            puntos.append((math.log(t), math.log(-math.log(1 - mr))))

    # Regresión Y sobre X:  ln(-ln(1-F)) = beta*ln(t) - beta*ln(eta)
    xs, ys = zip(*puntos)
    k = len(xs)
    mx, my = sum(xs) / k, sum(ys) / k
    sxx = sum((x - mx) ** 2 for x in xs)
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    syy = sum((y - my) ** 2 for y in ys)
    beta = sxy / sxx
    eta = math.exp(mx - my / beta)
    r2 = (sxy ** 2) / (sxx * syy) if syy > 0 else 1.0

    mttf = eta * math.gamma(1 + 1 / beta)
    b10 = eta * (-math.log(0.9)) ** (1 / beta)
    if beta < 0.95:
        patron = "Mortalidad infantil (tasa de falla decreciente)"
    elif beta <= 1.05:
        patron = "Fallas aleatorias (tasa de falla constante)"
    elif beta < 4:
        patron = "Desgaste (tasa de falla creciente)"
    else:
        patron = "Desgaste rápido / fin de vida abrupto"

    return {
        "beta_forma": round(beta, 3),
        "eta_escala": round(eta, 2),
        "R2_ajuste": round(r2, 4),
        "MTTF": round(mttf, 2),
        "vida_B10": round(b10, 2),
        "patron_falla": patron,
        "numero_fallas": len(tiempos_falla),
        "numero_suspensiones": len(tiempos_suspension),
        "advertencia": "Muestra pequeña: baja confianza en los parámetros." if len(tiempos_falla) < 6 else None,
    }


def confiabilidad_weibull(beta: float, eta: float, tiempo: float, gamma: float = 0.0) -> dict:
    """R(t), F(t) y tasa de riesgo h(t) para una distribución Weibull (2 o 3 parámetros)."""
    if beta <= 0 or eta <= 0:
        return {"error": "beta y eta deben ser positivos."}
    te = max(tiempo - gamma, 0.0)
    r = math.exp(-((te / eta) ** beta))
    h = (beta / eta) * (te / eta) ** (beta - 1) if te > 0 else (math.inf if beta < 1 else (1 / eta if beta == 1 else 0.0))
    return {
        "tiempo": tiempo,
        "R_confiabilidad_%": round(r * 100, 3),
        "F_prob_falla_%": round((1 - r) * 100, 3),
        "tasa_riesgo_h(t)": round(h, 8) if math.isfinite(h) else "infinita",
    }


def _integral_r(beta: float, eta: float, t: float, pasos: int = 400) -> float:
    """∫0^t R(x) dx por regla de Simpson."""
    if t <= 0:
        return 0.0
    pasos += pasos % 2
    dx = t / pasos
    f = lambda x: math.exp(-((x / eta) ** beta))
    s = f(0) + f(t) + sum((4 if i % 2 else 2) * f(i * dx) for i in range(1, pasos))
    return s * dx / 3


def intervalo_optimo_reemplazo(beta: float, eta: float, costo_preventivo: float,
                               costo_correctivo: float) -> dict:
    """Edad óptima de reemplazo preventivo que minimiza el costo por unidad de tiempo:
       C(T) = [Cp·R(T) + Cf·F(T)] / ∫0^T R(t) dt"""
    if costo_correctivo <= costo_preventivo:
        return {"recomendacion": "El correctivo no es más caro que el preventivo: no conviene el reemplazo preventivo por tiempo."}
    if beta <= 1:
        return {"recomendacion": "beta ≤ 1: la tasa de falla no aumenta con la edad, el reemplazo preventivo por tiempo "
                                 "no reduce fallas. Evaluar mantenimiento basado en condición (CBM) o rediseño."}
    costo_run_to_failure = costo_correctivo / (eta * math.gamma(1 + 1 / beta))
    mejor_t, mejor_c = None, math.inf
    t_max = eta * 3
    for i in range(1, 601):
        t = t_max * i / 600
        r = math.exp(-((t / eta) ** beta))
        c = (costo_preventivo * r + costo_correctivo * (1 - r)) / _integral_r(beta, eta, t)
        if c < mejor_c:
            mejor_t, mejor_c = t, c
    return {
        "intervalo_optimo": round(mejor_t, 1),
        "costo_por_unidad_tiempo_optimo": round(mejor_c, 4),
        "costo_por_unidad_tiempo_operar_hasta_falla": round(costo_run_to_failure, 4),
        "ahorro_%": round((1 - mejor_c / costo_run_to_failure) * 100, 2),
        "confiabilidad_al_intervalo_%": round(math.exp(-((mejor_t / eta) ** beta)) * 100, 2),
    }


def calcular_rpn(modos_falla: list[dict]) -> dict:
    """FMEA: RPN = Severidad × Ocurrencia × Detección (escalas 1–10), ordenado de mayor a menor."""
    salida = []
    for m in modos_falla:
        s, o, d = int(m.get("severidad", 0)), int(m.get("ocurrencia", 0)), int(m.get("deteccion", 0))
        if not all(1 <= v <= 10 for v in (s, o, d)):
            return {"error": f"Valores fuera de rango 1–10 en el modo '{m.get('modo', '?')}'."}
        rpn = s * o * d
        if s >= 9 or rpn >= 200:
            prioridad = "ALTA – acción inmediata"
        elif rpn >= 100:
            prioridad = "MEDIA – planificar acción"
        else:
            prioridad = "BAJA – monitorear"
        salida.append({"modo": m.get("modo", "sin nombre"), "S": s, "O": o, "D": d, "RPN": rpn, "prioridad": prioridad})
    salida.sort(key=lambda x: (x["RPN"], x["S"]), reverse=True)
    return {"ranking": salida}


def confiabilidad_sistema(configuracion: str, confiabilidades: list[float], k: int | None = None) -> dict:
    """Confiabilidad de un sistema: 'serie', 'paralelo' o 'k_de_n' (componentes idénticos o distintos)."""
    rs = confiabilidades
    if not rs or any(not 0 <= r <= 1 for r in rs):
        return {"error": "Las confiabilidades deben estar entre 0 y 1."}
    if configuracion == "serie":
        total = math.prod(rs)
    elif configuracion == "paralelo":
        total = 1 - math.prod(1 - r for r in rs)
    elif configuracion == "k_de_n":
        n = len(rs)
        if not k or not 1 <= k <= n:
            return {"error": "Para 'k_de_n' indica k entre 1 y n."}
        # Programación dinámica: probabilidad de que funcionen al menos k componentes
        dist = [1.0] + [0.0] * n
        for r in rs:
            dist = [dist[j] * (1 - r) + (dist[j - 1] * r if j > 0 else 0) for j in range(n + 1)]
        total = sum(dist[k:])
    else:
        return {"error": "Configuración no válida. Usa 'serie', 'paralelo' o 'k_de_n'."}
    return {"configuracion": configuracion, "confiabilidad_sistema_%": round(total * 100, 4)}


ISO14224_NIVELES = {
    1: "Industria", 2: "Categoría de negocio", 3: "Instalación", 4: "Planta / Unidad",
    5: "Sección / Sistema", 6: "Unidad de equipo", 7: "Subunidad",
    8: "Componente / Ítem mantenible", 9: "Parte",
}
ISO14224_CLASES = {
    "PU": ("Bombas", ["Transmisión de potencia", "Unidad de bomba", "Control y monitoreo", "Sistema de lubricación", "Misceláneos"]),
    "CO": ("Compresores", ["Transmisión de potencia", "Compresor", "Control y monitoreo", "Sistema de lubricación", "Sistema de sellado del eje", "Misceláneos"]),
    "EM": ("Motores eléctricos", ["Motor eléctrico", "Control y monitoreo", "Sistema de lubricación", "Sistema de enfriamiento", "Misceláneos"]),
    "EG": ("Generadores eléctricos", ["Transmisión de potencia", "Generador", "Control y monitoreo", "Sistema de lubricación", "Sistema de enfriamiento", "Misceláneos"]),
    "GT": ("Turbinas de gas", ["Sistema de arranque", "Entrada de aire", "Generador de gas", "Turbina de potencia", "Control y monitoreo", "Sistema de lubricación", "Sistema de combustible", "Misceláneos"]),
    "ST": ("Turbinas de vapor", ["Transmisión de potencia", "Turbina de vapor", "Control y monitoreo", "Sistema de lubricación", "Condensador", "Misceláneos"]),
    "CE": ("Motores de combustión", ["Sistema de arranque", "Motor de combustión", "Control y monitoreo", "Sistema de lubricación", "Sistema de enfriamiento", "Misceláneos"]),
    "HE": ("Intercambiadores de calor", ["Externos", "Internos", "Control y monitoreo", "Misceláneos"]),
    "VE": ("Recipientes a presión", ["Recipiente", "Control y monitoreo", "Misceláneos"]),
    "VA": ("Válvulas", ["Válvula", "Actuador", "Control y monitoreo", "Misceláneos"]),
    "PI": ("Tuberías", ["Tubería", "Válvulas", "Control y monitoreo", "Misceláneos"]),
}


def taxonomia_iso14224(codigo_clase: str | None = None) -> dict:
    """Devuelve los 9 niveles de ISO 14224 y, si se indica, las subunidades típicas de una clase de equipo."""
    res = {"niveles": ISO14224_NIVELES,
           "clases_disponibles": {k: v[0] for k, v in ISO14224_CLASES.items()}}
    if codigo_clase:
        c = codigo_clase.upper()
        if c in ISO14224_CLASES:
            res["clase"] = {"codigo": c, "nombre": ISO14224_CLASES[c][0], "subunidades_tipicas": ISO14224_CLASES[c][1]}
        else:
            res["aviso"] = f"Clase '{codigo_clase}' no está en el catálogo resumido."
    return res


# ---------------------------------------------------------------------------
# Definición de herramientas para Claude
# ---------------------------------------------------------------------------
_num_list = {"type": "array", "items": {"type": "number"}}

HERRAMIENTAS = [
    {
        "name": "calcular_indicadores",
        "description": "Calcula MTBF, MTTR, disponibilidad inherente y tasa de falla a partir de tiempos entre fallas y tiempos de reparación (en horas).",
        "input_schema": {"type": "object", "properties": {
            "tiempos_entre_fallas": {**_num_list, "description": "Tiempos de operación entre fallas (h)."},
            "tiempos_reparacion": {**_num_list, "description": "Duración de cada reparación (h)."},
            "horas_periodo": {"type": "number", "description": "Opcional: horas calendario del periodo para disponibilidad operacional."}},
            "required": ["tiempos_entre_fallas", "tiempos_reparacion"]},
    },
    {
        "name": "ajuste_weibull",
        "description": "Ajusta una distribución Weibull de 2 parámetros (beta, eta) a tiempos de falla, admitiendo suspensiones (equipos retirados sin fallar). Devuelve beta, eta, R², MTTF, B10 y patrón de falla.",
        "input_schema": {"type": "object", "properties": {
            "tiempos_falla": {**_num_list, "description": "Tiempos hasta la falla."},
            "tiempos_suspension": {**_num_list, "description": "Opcional: tiempos de unidades censuradas (sin falla)."}},
            "required": ["tiempos_falla"]},
    },
    {
        "name": "confiabilidad_weibull",
        "description": "Calcula R(t), F(t) y la tasa de riesgo h(t) en un tiempo dado para una Weibull con parámetros conocidos.",
        "input_schema": {"type": "object", "properties": {
            "beta": {"type": "number"}, "eta": {"type": "number"}, "tiempo": {"type": "number"},
            "gamma": {"type": "number", "description": "Opcional: parámetro de localización (vida mínima)."}},
            "required": ["beta", "eta", "tiempo"]},
    },
    {
        "name": "intervalo_optimo_reemplazo",
        "description": "Calcula el intervalo óptimo de reemplazo preventivo por edad que minimiza el costo por unidad de tiempo, dados beta, eta y los costos de una intervención preventiva y una correctiva.",
        "input_schema": {"type": "object", "properties": {
            "beta": {"type": "number"}, "eta": {"type": "number"},
            "costo_preventivo": {"type": "number"}, "costo_correctivo": {"type": "number"}},
            "required": ["beta", "eta", "costo_preventivo", "costo_correctivo"]},
    },
    {
        "name": "calcular_rpn",
        "description": "FMEA: calcula el RPN (S×O×D, escalas 1–10) de cada modo de falla y los ordena por prioridad.",
        "input_schema": {"type": "object", "properties": {
            "modos_falla": {"type": "array", "items": {"type": "object", "properties": {
                "modo": {"type": "string"}, "severidad": {"type": "integer"},
                "ocurrencia": {"type": "integer"}, "deteccion": {"type": "integer"}},
                "required": ["modo", "severidad", "ocurrencia", "deteccion"]}}},
            "required": ["modos_falla"]},
    },
    {
        "name": "confiabilidad_sistema",
        "description": "Confiabilidad de un sistema de componentes en serie, en paralelo (redundancia activa) o k-de-n.",
        "input_schema": {"type": "object", "properties": {
            "configuracion": {"type": "string", "enum": ["serie", "paralelo", "k_de_n"]},
            "confiabilidades": {**_num_list, "description": "Confiabilidad de cada componente (0 a 1)."},
            "k": {"type": "integer", "description": "Solo para k_de_n: mínimo de componentes que deben funcionar."}},
            "required": ["configuracion", "confiabilidades"]},
    },
    {
        "name": "taxonomia_iso14224",
        "description": "Consulta los 9 niveles jerárquicos de ISO 14224 y las subunidades típicas de una clase de equipo (PU, CO, EM, EG, GT, ST, CE, HE, VE, VA, PI).",
        "input_schema": {"type": "object", "properties": {
            "codigo_clase": {"type": "string", "description": "Opcional: código de clase de equipo, p. ej. 'PU'."}}},
    },
]

FUNCIONES = {
    "calcular_indicadores": calcular_indicadores,
    "ajuste_weibull": ajuste_weibull,
    "confiabilidad_weibull": confiabilidad_weibull,
    "intervalo_optimo_reemplazo": intervalo_optimo_reemplazo,
    "calcular_rpn": calcular_rpn,
    "confiabilidad_sistema": confiabilidad_sistema,
    "taxonomia_iso14224": taxonomia_iso14224,
}



MODELO = os.environ.get("MODELO_OLLAMA", "qwen2.5:14b")
TEMPERATURA = 0.2
CONTEXTO = 8192
URL_OLLAMA = os.environ.get("OLLAMA_HOST", "http://localhost:11434").rstrip("/") + "/api/chat"
PROVEEDOR = "ollama"            # "ollama" (tu PC) o "gemini" (nube)
GEMINI_API_KEY = ""
GEMINI_URL = "https://generativelanguage.googleapis.com/v1beta/openai/chat/completions"
MAX_PASOS_HERRAMIENTAS = 8

# Convierte las herramientas al formato que entiende Ollama
HERRAMIENTAS_OLLAMA = [
    {"type": "function",
     "function": {"name": h["name"], "description": h["description"], "parameters": h["input_schema"]}}
    for h in HERRAMIENTAS
]


def llamar_ollama(mensajes: list[dict]) -> dict:
    datos = json.dumps({
        "model": MODELO,
        "messages": mensajes,
        "tools": HERRAMIENTAS_OLLAMA,
        "stream": False,
        "options": {"temperature": TEMPERATURA, "num_ctx": CONTEXTO},
    }).encode("utf-8")
    req = urllib.request.Request(URL_OLLAMA, data=datos, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=600) as resp:
        return json.loads(resp.read().decode("utf-8"))["message"]


MODELOS_RESPALDO = ["gemini-3.5-flash", "gemini-2.5-flash"]   # se usan si el principal está saturado
_ultimo_modelo_ok = None


def _pedir_gemini(modelo: str, mensajes: list[dict]) -> dict:
    datos = json.dumps({"model": modelo, "messages": mensajes, "tools": HERRAMIENTAS_OLLAMA}).encode("utf-8")
    req = urllib.request.Request(GEMINI_URL, data=datos, headers={
        "Content-Type": "application/json", "Authorization": f"Bearer {GEMINI_API_KEY}"})
    with urllib.request.urlopen(req, timeout=300) as resp:
        return json.loads(resp.read().decode("utf-8"))["choices"][0]["message"]


def llamar_gemini(mensajes: list[dict]) -> dict:
    """Gemini (API compatible con OpenAI) con reintentos automáticos y modelos de respaldo.
    - 500/502/503/504 (servidor saturado): reintenta con espera creciente y luego pasa al siguiente modelo.
    - 429 (límite por minuto): espera un momento, reintenta y luego prueba otro modelo (tiene su propio cupo).
    - 404 (modelo no existe): pasa directo al siguiente modelo."""
    global _ultimo_modelo_ok
    import time
    limpios = [{k: v for k, v in m.items() if k != "tool_name"} for m in mensajes]
    orden = [MODELO] + [m for m in MODELOS_RESPALDO if m != MODELO]
    if _ultimo_modelo_ok in orden:                 # dentro de una misma pregunta, sigue con el que funcionó
        orden.remove(_ultimo_modelo_ok)
        orden.insert(0, _ultimo_modelo_ok)
    ultimo_error = None
    for modelo in orden:
        for intento in range(3):
            try:
                mensaje = _pedir_gemini(modelo, limpios)
                _ultimo_modelo_ok = modelo
                return mensaje
            except urllib.error.HTTPError as e:
                ultimo_error = e
                if e.code in (500, 502, 503, 504):
                    time.sleep(1.5 * (intento + 1))
                    continue
                if e.code == 429 and intento == 0:
                    time.sleep(3)
                    continue
                if e.code in (404, 429):
                    break                          # probar el siguiente modelo
                raise
            except urllib.error.URLError as e:
                ultimo_error = e
                time.sleep(1.5)
    raise ultimo_error


def llamar_modelo(mensajes: list[dict]) -> dict:
    return llamar_gemini(mensajes) if PROVEEDOR == "gemini" else llamar_ollama(mensajes)


class AgenteLocal:
    def __init__(self):
        self.mensajes = [{"role": "system", "content": SYSTEM_PROMPT}]

    def _ejecutar(self, nombre: str, argumentos) -> str:
        if isinstance(argumentos, str):          # algunos modelos envían los argumentos como texto JSON
            try:
                argumentos = json.loads(argumentos)
            except json.JSONDecodeError:
                argumentos = {}
        funcion = FUNCIONES.get(nombre)
        if funcion is None:
            resultado = {"error": f"Herramienta desconocida: {nombre}"}
        else:
            try:
                resultado = funcion(**argumentos)
            except Exception as e:
                resultado = {"error": f"{type(e).__name__}: {e}"}
        print(f"  ⚙ {nombre} → {json.dumps(resultado, ensure_ascii=False)[:300]}")
        return json.dumps(resultado, ensure_ascii=False)

    def preguntar(self, texto: str) -> str:
        self.mensajes.append({"role": "user", "content": texto})
        for _ in range(MAX_PASOS_HERRAMIENTAS):
            mensaje = llamar_modelo(self.mensajes)
            self.mensajes.append(mensaje)
            llamadas = mensaje.get("tool_calls") or []
            if not llamadas:
                return mensaje.get("content", "")
            for llamada in llamadas:
                f = llamada["function"]
                respuesta_herramienta = {"role": "tool", "tool_name": f["name"],
                                         "content": self._ejecutar(f["name"], f.get("arguments", {}))}
                if llamada.get("id"):
                    respuesta_herramienta["tool_call_id"] = llamada["id"]
                self.mensajes.append(respuesta_herramienta)
        return "(Se alcanzó el límite de pasos con herramientas. Intenta reformular la pregunta.)"

    def reiniciar(self):
        self.mensajes = self.mensajes[:1]



# =============================================================================
# BASE DE CONOCIMIENTO PROPIA (tus documentos: PDF, Word, Excel, CSV, TXT, MD)
# =============================================================================
import re
import unicodedata
from pathlib import Path

CARPETA = Path(__file__).resolve().parent
CARPETA_CONOCIMIENTO = CARPETA / "conocimiento"
ARCHIVO_INDICE = CARPETA / "conocimiento_indice.json"
EXTENSIONES = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}
LIBRERIAS = "python -m pip install pypdf python-docx openpyxl"

_STOP = set("""a al algo ante antes como con contra cual cuando de del desde donde e el ella en entre era es esa ese
eso esta este esto estos estas fue ha han hay la las le les lo los mas me mi mis muy nada ni no nos o otra otro para
pero poco por porque que quien se sea segun ser si sin sobre son su sus tambien te tiene tu tus un una uno unos unas
y ya cada cuanto cuantos cuantas debe deben puede hace hacer cual cuales the of and to in is for on with by""".split())


def _normalizar(texto: str) -> list[str]:
    texto = unicodedata.normalize("NFD", texto.lower())
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return [t for t in re.findall(r"[a-z0-9]+", texto) if len(t) > 1 and t not in _STOP]


def extraer_texto(ruta: Path) -> list[tuple[str, str]]:
    """Devuelve [(ubicación, texto)] del archivo."""
    ext = ruta.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return [(f"p. {i}", p.extract_text() or "") for i, p in enumerate(PdfReader(str(ruta)).pages, 1)]
    if ext == ".docx":
        import docx
        d = docx.Document(str(ruta))
        partes = [p.text for p in d.paragraphs if p.text.strip()]
        for tabla in d.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text.strip() for c in fila.cells))
        return [("documento", "\n".join(partes))]
    if ext == ".xlsx":
        import openpyxl
        libro = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
        salida = []
        for hoja in libro.worksheets:
            filas = [" | ".join(str(c) for c in fila if c is not None) for fila in hoja.iter_rows(values_only=True)]
            salida.append((f"hoja {hoja.title}", "\n".join(f for f in filas if f.strip())))
        libro.close()
        return salida
    return [("documento", ruta.read_text(encoding="utf-8", errors="ignore"))]


def fragmentar(texto: str, tam: int = 900, solape: int = 150) -> list[str]:
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()
    salida, i = [], 0
    while i < len(texto):
        fin = min(i + tam, len(texto))
        if fin < len(texto):
            corte = max(texto.rfind("\n", i + tam // 2, fin), texto.rfind(". ", i + tam // 2, fin))
            if corte <= 0:
                corte = texto.rfind(" ", i + tam // 2, fin)
            fin = corte + 1 if corte > 0 else fin
        frag = texto[i:fin].strip()
        if len(frag) > 15:
            salida.append(frag)
        if fin >= len(texto):
            break
        i = max(fin - solape, i + 1)
    return salida


def embeber(textos: list[str], servidor: str, modelo: str, es_consulta: bool = False):
    """Vectores de significado con Ollama. Devuelve None si el modelo no está disponible."""
    if not modelo:
        return None
    if "nomic" in modelo:
        prefijo = "search_query: " if es_consulta else "search_document: "
        textos = [prefijo + t for t in textos]
    vectores = []
    try:
        for i in range(0, len(textos), 16):
            datos = json.dumps({"model": modelo, "input": textos[i:i + 16]}).encode("utf-8")
            req = urllib.request.Request(servidor.rstrip("/") + "/api/embed", data=datos,
                                         headers={"Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=300) as r:
                vectores += json.loads(r.read().decode("utf-8"))["embeddings"]
    except Exception:
        return None
    return [[round(x, 5) for x in v] for v in vectores]


def _sin_vector(f: dict, puntaje: float) -> dict:
    return {"doc": f["doc"], "loc": f["loc"], "texto": f["texto"], "puntaje": puntaje}


class BaseConocimiento:
    def __init__(self):
        self.archivos, self.fragmentos = {}, []
        try:
            datos = json.loads(ARCHIVO_INDICE.read_text(encoding="utf-8"))
            self.archivos, self.fragmentos = datos["archivos"], datos["fragmentos"]
        except (OSError, ValueError, KeyError):
            pass
        self._tokens = None

    def _guardar(self):
        ARCHIVO_INDICE.write_text(json.dumps({"archivos": self.archivos, "fragmentos": self.fragmentos},
                                             ensure_ascii=False), encoding="utf-8")

    @property
    def busqueda_inteligente(self) -> bool:
        return bool(self.fragmentos) and all("emb" in f for f in self.fragmentos)

    def actualizar(self, servidor: str, modelo_emb: str, progreso=None) -> list[str]:
        """Procesa archivos nuevos o modificados de la carpeta y quita los borrados. Devuelve avisos."""
        CARPETA_CONOCIMIENTO.mkdir(exist_ok=True)
        actuales = {p.name: p for p in CARPETA_CONOCIMIENTO.iterdir()
                    if p.is_file() and p.suffix.lower() in EXTENSIONES and not p.name.startswith("~$")}
        firma = lambda p: [int(p.stat().st_mtime), p.stat().st_size]
        cambiados = [n for n, p in actuales.items()
                     if n not in self.archivos or self.archivos[n].get("firma") != firma(p)
                     or self.archivos[n].get("error")]          # reintenta los que fallaron
        quitar = set(self.archivos) - set(actuales) | set(cambiados)
        self.fragmentos = [f for f in self.fragmentos if f["doc"] not in quitar]
        for n in quitar:
            self.archivos.pop(n, None)

        avisos = []
        for k, nombre in enumerate(sorted(cambiados), 1):
            if progreso:
                progreso(k / len(cambiados), f"Procesando {nombre} ({k}/{len(cambiados)})")
            ruta = actuales[nombre]
            info = {"firma": firma(ruta), "fragmentos": 0, "error": None}
            try:
                nuevos = [{"doc": nombre, "loc": loc, "texto": frag}
                          for loc, texto in extraer_texto(ruta) for frag in fragmentar(texto)]
                if not nuevos:
                    info["error"] = "Sin texto legible (¿PDF escaneado como imagen?)"
                else:
                    vectores = embeber([f["texto"] for f in nuevos], servidor, modelo_emb)
                    if vectores and len(vectores) == len(nuevos):
                        for f, v in zip(nuevos, vectores):
                            f["emb"] = v
                    self.fragmentos += nuevos
                    info["fragmentos"] = len(nuevos)
            except ImportError as e:
                info["error"] = f"Falta la librería '{e.name}'. En la terminal ejecuta: {LIBRERIAS}"
            except Exception as e:
                info["error"] = f"No se pudo leer: {type(e).__name__}: {e}"
            if info["error"]:
                avisos.append(f"{nombre}: {info['error']}")
            self.archivos[nombre] = info

        # si el modelo de búsqueda ya está disponible, completa los vectores que falten
        faltan = [f for f in self.fragmentos if "emb" not in f]
        if faltan:
            vectores = embeber([f["texto"] for f in faltan], servidor, modelo_emb)
            if vectores and len(vectores) == len(faltan):
                for f, v in zip(faltan, vectores):
                    f["emb"] = v
        self._tokens = None
        self._guardar()
        return avisos

    def _por_significado(self, consulta: str, servidor: str, modelo_emb: str):
        if not self.busqueda_inteligente:
            return None
        q = embeber([consulta], servidor, modelo_emb, es_consulta=True)
        if not q or len(q[0]) != len(self.fragmentos[0]["emb"]):
            return None
        q = q[0]
        nq = math.sqrt(sum(x * x for x in q)) or 1.0
        puntuados = []
        for i, f in enumerate(self.fragmentos):
            v = f["emb"]
            cos = sum(a * b for a, b in zip(q, v)) / (nq * (math.sqrt(sum(x * x for x in v)) or 1.0))
            if cos >= 0.3:
                puntuados.append((cos, i))
        return sorted(puntuados, reverse=True)

    def _por_palabras(self, consulta: str):
        """BM25: ideal para códigos de equipos, normas y términos exactos."""
        if self._tokens is None:
            self._tokens = [_normalizar(f["doc"] + " " + f["texto"]) for f in self.fragmentos]
        terminos = set(_normalizar(consulta))
        if not terminos:
            return []
        n = len(self._tokens)
        promedio = sum(len(t) for t in self._tokens) / n or 1.0
        df = {t: sum(1 for toks in self._tokens if t in toks) for t in terminos}
        puntuados = []
        for i, toks in enumerate(self._tokens):
            s = 0.0
            for t in terminos:
                tf = toks.count(t)
                if tf:
                    idf = math.log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
                    s += idf * tf * 2.5 / (tf + 1.5 * (0.25 + 0.75 * len(toks) / promedio))
            if s > 0:
                puntuados.append((s, i))
        return sorted(puntuados, reverse=True)

    def buscar(self, consulta: str, k: int, servidor: str, modelo_emb: str) -> list[dict]:
        if not self.fragmentos:
            return []
        palabras = self._por_palabras(consulta)
        significado = self._por_significado(consulta, servidor, modelo_emb)
        if significado is None:
            return [_sin_vector(self.fragmentos[i], round(s, 2)) for s, i in palabras[:k]]
        # búsqueda híbrida: fusión de rangos recíprocos (significado + palabras clave)
        fusion = {}
        for lista in (significado, palabras):
            for rango, (_, i) in enumerate(lista[:30]):
                fusion[i] = fusion.get(i, 0.0) + 1.0 / (60 + rango)
        mejores = sorted(fusion.items(), key=lambda x: x[1], reverse=True)[:k]
        return [_sin_vector(self.fragmentos[i], round(s * 100, 2)) for i, s in mejores]

# =============================================================================
# INTERFAZ WEB PROFESIONAL (Streamlit)
# =============================================================================
import base64
import urllib.request
from datetime import datetime
from pathlib import Path

import streamlit as st

CARPETA = Path(__file__).resolve().parent
ARCHIVO_AJUSTES = CARPETA / "confiabot_ajustes.json"
ARCHIVO_LOGO = CARPETA / "confiabot_logo_empresa.png"

AJUSTES_DEFECTO = {
    "nombre_asistente": "Agente IA de Aldair",
    "nombre_empresa": "",
    "color_principal": "#1F3864",
    "color_acento": "#F4B183",
    "modelo": "qwen2.5:14b",
    "temperatura": 0.2,
    "servidor": "http://localhost:11434",
    "nivel_detalle": "Normal",
    "unidad_tiempo": "horas",
    "mostrar_calculos": True,
    "instrucciones_extra": "",
    "contexto": 8192,
    "usar_conocimiento": True,
    "fragmentos": 4,
    "modelo_embeddings": "nomic-embed-text",
    "modelo_gemini": "gemini-3.8-flash",
}


def secreto(nombre: str) -> str:
    """Lee un secreto de Streamlit Cloud (o variable de entorno). Vacío si no existe."""
    try:
        valor = st.secrets.get(nombre, "")
    except Exception:
        valor = ""
    return str(valor or os.environ.get(nombre, ""))

NOMBRES_HERRAMIENTAS = {
    "calcular_indicadores": "Indicadores MTBF · MTTR · Disponibilidad",
    "ajuste_weibull": "Análisis de Weibull",
    "confiabilidad_weibull": "Confiabilidad R(t)",
    "intervalo_optimo_reemplazo": "Intervalo óptimo de reemplazo",
    "calcular_rpn": "FMEA · Ranking RPN",
    "confiabilidad_sistema": "Confiabilidad de sistema",
    "taxonomia_iso14224": "Taxonomía ISO 14224",
}

EJEMPLOS = [
    ("📈 Análisis Weibull", "Tiempos de falla de rodamientos: 1200, 1850, 2100, 2600, 3100, 3900 h. Haz un análisis Weibull e interprétalo."),
    ("🔧 Reemplazo óptimo", "Beta 2.8, eta 3000 h, preventivo 800 USD y correctivo 6000 USD. ¿Cada cuánto debo reemplazar el componente?"),
    ("⏱️ MTBF y disponibilidad", "Tiempos entre fallas: 500, 700, 650 h. Reparaciones: 5, 8, 6 h. Calcula MTBF, MTTR y disponibilidad."),
    ("⚠️ FMEA con RPN", "Arma un FMEA con RPN para una bomba centrífuga de crudo y prioriza las acciones."),
    ("🔗 Confiabilidad de sistema", "Tengo 3 bombas en paralelo con confiabilidad de 0.9 cada una y necesito al menos 2 funcionando. ¿Cuál es la confiabilidad del sistema?"),
    ("🌳 Taxonomía ISO 14224", "Dame los niveles de ISO 14224 y las subunidades de un compresor."),
]


# ---------------------------------------------------------------- utilidades
def cargar_ajustes() -> dict:
    ajustes = dict(AJUSTES_DEFECTO)
    try:
        ajustes.update(json.loads(ARCHIVO_AJUSTES.read_text(encoding="utf-8")))
    except (OSError, ValueError):
        pass
    if ajustes["nombre_asistente"] == "ConfiaBot":          # nombre de la versión anterior
        ajustes["nombre_asistente"] = AJUSTES_DEFECTO["nombre_asistente"]
    return ajustes


def guardar_ajustes(ajustes: dict) -> None:
    ARCHIVO_AJUSTES.write_text(json.dumps(ajustes, ensure_ascii=False, indent=2), encoding="utf-8")


def logo_svg(principal: str, acento: str) -> str:
    return f"""<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64">
<defs><linearGradient id="g" x1="0" y1="0" x2="1" y2="1">
<stop offset="0" stop-color="{principal}"/><stop offset="1" stop-color="#2E86AB"/></linearGradient></defs>
<polygon points="32,3 57,17.5 57,46.5 32,61 7,46.5 7,17.5" fill="url(#g)"/>
<circle cx="32" cy="32" r="16" fill="none" stroke="#ffffff" stroke-opacity=".25" stroke-width="3"/>
<polyline points="12,34 21,34 25.5,23 31,43 36,27 40,34 52,34" fill="none" stroke="{acento}"
 stroke-width="3.4" stroke-linecap="round" stroke-linejoin="round"/></svg>"""


def img_svg(svg: str, tam: int) -> str:
    b64 = base64.b64encode(svg.encode("utf-8")).decode()
    return f'<img src="data:image/svg+xml;base64,{b64}" width="{tam}" height="{tam}" alt="logo">'


def img_logo_empresa(alto: int) -> str:
    if not ARCHIVO_LOGO.exists():
        return ""
    b64 = base64.b64encode(ARCHIVO_LOGO.read_bytes()).decode()
    return f'<img class="cb-logo-empresa" src="data:image/png;base64,{b64}" style="max-height:{alto}px" alt="empresa">'


@st.cache_data(ttl=20, show_spinner=False)
def estado_ollama(servidor: str):
    """(conectado, [modelos instalados])"""
    try:
        with urllib.request.urlopen(servidor.rstrip("/") + "/api/tags", timeout=3) as r:
            datos = json.loads(r.read().decode("utf-8"))
        return True, sorted(m["name"] for m in datos.get("models", []))
    except Exception:
        return False, []


def prompt_sistema(aj: dict) -> str:
    detalle = {
        "Breve": "Responde de forma breve y directa: resultados clave y una recomendación, máximo 150 palabras.",
        "Normal": "Responde con un nivel de detalle equilibrado.",
        "Detallado": "Responde de forma detallada: método, supuestos, resultados, interpretación y plan de acción.",
    }[aj["nivel_detalle"]]
    extra = [f"\nTu nombre es {aj['nombre_asistente']}. Eres un agente de IA propio y personal, "
             "creado y entrenado por tu usuario con su propio conocimiento.",
             "Cuando el mensaje traiga 'Información de mi base de conocimiento', úsala con prioridad sobre tu "
             "conocimiento general y cita la fuente así: [nombre del documento, página]. Si esa información no "
             "responde la pregunta, dilo y responde con tu conocimiento general. Nunca inventes citas."]
    if aj["nombre_empresa"]:
        extra.append(f"Trabajas para la empresa {aj['nombre_empresa']}.")
    extra.append(detalle)
    extra.append(f"Si el usuario no indica unidades de tiempo, asume que están en {aj['unidad_tiempo']}.")
    if aj["instrucciones_extra"].strip():
        extra.append("\nInstrucciones y conocimiento propios de la empresa (tienen prioridad):\n"
                     + aj["instrucciones_extra"].strip())
    return SYSTEM_PROMPT + "\n".join(extra)


class AgenteWeb(AgenteLocal):
    """Agente local que además guarda los cálculos para mostrarlos en la página."""
    def __init__(self):
        super().__init__()
        self.registro = []

    def _ejecutar(self, nombre, argumentos):
        resultado = super()._ejecutar(nombre, argumentos)
        self.registro.append((nombre, argumentos, resultado))
        return resultado

    def preguntar_con_conocimiento(self, pregunta: str, fuentes: list[dict]) -> str:
        if not fuentes:
            return self.preguntar(pregunta)
        contexto = "\n\n".join(f"[{i}] {f['doc']} ({f['loc']}):\n{f['texto']}" for i, f in enumerate(fuentes, 1))
        texto = (f"Información de mi base de conocimiento (úsala si es relevante y cita el documento):\n"
                 f"{contexto}\n\nMi pregunta: {pregunta}")
        indice = len(self.mensajes)
        try:
            return self.preguntar(texto)
        finally:
            # en el historial queda solo la pregunta, para no llenar la memoria del modelo
            if len(self.mensajes) > indice and self.mensajes[indice].get("role") == "user":
                docs = ", ".join(sorted({f["doc"] for f in fuentes}))
                self.mensajes[indice] = {"role": "user", "content": f"{pregunta}\n[Documentos consultados: {docs}]"}


def mostrar_fuentes(fuentes):
    if not fuentes:
        return
    with st.expander(f"📚 Fuentes de tu base de conocimiento ({len(fuentes)})"):
        for i, f in enumerate(fuentes, 1):
            st.markdown(f"**[{i}] {f['doc']}** · {f['loc']}")
            st.caption(f["texto"][:350] + ("…" if len(f["texto"]) > 350 else ""))


def mostrar_calculos(calculos):
    for nombre, args, res in calculos:
        titulo = NOMBRES_HERRAMIENTAS.get(nombre, nombre)
        with st.expander(f"🧮 {titulo}"):
            try:
                args = json.loads(args) if isinstance(args, str) else args
                res = json.loads(res)
            except ValueError:
                st.code(str(res))
                continue
            c1, c2 = st.columns(2)
            with c1:
                st.markdown("**Datos de entrada**")
                st.json(args, expanded=False)
            with c2:
                st.markdown("**Resultado**")
                if "ranking" in res:
                    st.dataframe(res["ranking"], hide_index=True, use_container_width=True)
                else:
                    filas = [{"Indicador": k.replace("_", " "), "Valor": str(v)}
                             for k, v in res.items() if v is not None and not isinstance(v, (dict, list))]
                    if filas:
                        st.dataframe(filas, hide_index=True, use_container_width=True)
                    anidados = {k: v for k, v in res.items() if isinstance(v, (dict, list))}
                    if anidados:
                        st.json(anidados, expanded=False)


def conversacion_markdown(chat, aj) -> str:
    lineas = [f"# Conversación con {aj['nombre_asistente']}",
              f"_{datetime.now():%Y-%m-%d %H:%M}_ · modelo {aj['modelo']}\n"]
    for rol, texto, *_ in chat:
        lineas.append(f"## {'Usuario' if rol == 'user' else aj['nombre_asistente']}\n\n{texto}\n")
    return "\n".join(lineas)


# ---------------------------------------------------------------- estado y ajustes
st.set_page_config(page_title=f"{cargar_ajustes()['nombre_asistente']} · Confiabilidad", page_icon="⚙️",
                   layout="wide", initial_sidebar_state="expanded")

if "ajustes" not in st.session_state:
    st.session_state.ajustes = cargar_ajustes()
if "agente" not in st.session_state:
    st.session_state.agente = AgenteWeb()
    st.session_state.chat = []
    st.session_state.n_calculos = 0
st.session_state.setdefault("pagina", "💬 Asistente")
st.session_state.setdefault("pendiente", None)
if "base" not in st.session_state:
    st.session_state.base = BaseConocimiento()
BASE = st.session_state.base

AJ = st.session_state.ajustes
GEMINI_API_KEY = secreto("GEMINI_API_KEY")
EN_NUBE = bool(GEMINI_API_KEY)                 # con clave de Gemini = versión en la nube
CLAVE_ADMIN = secreto("CLAVE_ADMIN")           # si existe, se activa el modo compañeros
st.session_state.setdefault("admin", not CLAVE_ADMIN)
ES_ADMIN = st.session_state.admin
TEMPERATURA = float(AJ["temperatura"])
CONTEXTO = int(AJ["contexto"])
URL_OLLAMA = AJ["servidor"].rstrip("/") + "/api/chat"
if EN_NUBE:
    PROVEEDOR = "gemini"
    MODELO = secreto("MODELO_GEMINI") or AJ["modelo_gemini"]
    MODELO_EMB = ""                             # en la nube: búsqueda por palabras clave
    CONECTADO, MODELOS = True, []
    if "indexado" not in st.session_state:      # procesa los documentos subidos al repositorio
        BASE.actualizar("", "")
        st.session_state.indexado = True
else:
    PROVEEDOR = "ollama"
    MODELO = AJ["modelo"]
    MODELO_EMB = AJ["modelo_embeddings"]
    CONECTADO, MODELOS = estado_ollama(AJ["servidor"])
st.session_state.agente.mensajes[0]["content"] = prompt_sistema(AJ)

# ---------------------------------------------------------------- estilos
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
html, body, .stApp, .stMarkdown, button, input, textarea, select { font-family: 'Inter', sans-serif !important; }
#MainMenu, footer, .stDeployButton, [data-testid="stAppDeployButton"], [data-testid="stToolbarActions"]
  { visibility: hidden; display: none; }
.block-container { padding-top: 1.6rem; max-width: 1100px; }
.cb-header { display:flex; align-items:center; flex-wrap:wrap; gap:16px; padding:18px 22px; border-radius:16px;
  background: linear-gradient(120deg, %PRINCIPAL% 0%, %PRINCIPAL% 55%, #2E86AB 100%); color:#fff; margin-bottom:20px;
  box-shadow: 0 6px 20px rgba(31,56,100,.18); }
.cb-header h1 { font-size:1.55rem; font-weight:700; margin:0; padding:0; color:#fff; line-height:1.25; }
.cb-sub { opacity:.85; font-size:.9rem; margin-top:2px; }
.cb-der { margin-left:auto; display:flex; align-items:center; gap:14px; flex-wrap:wrap; }
.cb-pill { background: rgba(255,255,255,.14); border:1px solid rgba(255,255,255,.25); padding:6px 12px;
  border-radius:999px; font-size:.8rem; white-space:nowrap; }
.cb-dot { display:inline-block; width:8px; height:8px; border-radius:50%; margin-right:6px; vertical-align:middle; }
.cb-logo-empresa { background:#fff; border-radius:10px; padding:6px 10px; }
.cb-marca { display:flex; align-items:center; gap:10px; margin-bottom:6px; }
.cb-marca b { font-size:1.15rem; color:%PRINCIPAL%; }
.cb-marca span { display:block; font-size:.75rem; color:#6b7280; }
.cb-bienvenida { text-align:center; padding:10px 0 4px; }
.cb-bienvenida h3 { color:%PRINCIPAL%; margin-bottom:4px; }
.cb-bienvenida p { color:#6b7280; margin-top:0; }
.stButton > button { border-radius:10px; border:1px solid #d9dee8; font-weight:500; transition: all .15s; }
.stButton > button:hover { border-color:%PRINCIPAL%; color:%PRINCIPAL%; transform: translateY(-1px); }
.stButton > button[kind="primary"], .stFormSubmitButton > button { background:%PRINCIPAL%; border-color:%PRINCIPAL%; color:#fff; }
.stButton > button[kind="primary"]:hover, .stFormSubmitButton > button:hover { color:#fff; filter:brightness(1.15); }
[data-testid="stChatMessage"] { border-radius:14px; padding:12px 16px; }
[data-testid="stSidebar"] { border-right:1px solid #e5e7eb; }
.cb-pie { font-size:.75rem; color:#9ca3af; margin-top:12px; }
@media (max-width: 640px) { .cb-header h1 { font-size:1.2rem; } .cb-der { margin-left:0; } }
</style>
"""
st.markdown(CSS.replace("%PRINCIPAL%", AJ["color_principal"]), unsafe_allow_html=True)

LOGO = logo_svg(AJ["color_principal"], AJ["color_acento"])
LOGO_HEADER = logo_svg("#FFFFFF00", AJ["color_acento"]).replace('fill="url(#g)"', 'fill="rgba(255,255,255,.12)"')

# ---------------------------------------------------------------- barra lateral
with st.sidebar:
    st.markdown(f"""<div class="cb-marca">{img_svg(LOGO, 44)}
        <div><b>{AJ['nombre_asistente']}</b><span>Agente IA propio · Confiabilidad</span></div></div>""",
                unsafe_allow_html=True)
    if ARCHIVO_LOGO.exists():
        st.markdown(img_logo_empresa(46), unsafe_allow_html=True)
    st.divider()
    paginas = (["💬 Asistente", "📚 Conocimiento", "⚙️ Ajustes", "ℹ️ Acerca de"] if ES_ADMIN
               else ["💬 Asistente", "ℹ️ Acerca de"])
    if st.session_state.pagina not in paginas:
        st.session_state.pagina = "💬 Asistente"
    st.radio("Navegación", paginas, key="pagina", label_visibility="collapsed")
    st.divider()
    if st.button("➕ Nueva conversación", use_container_width=True, type="primary"):
        st.session_state.agente = AgenteWeb()
        st.session_state.chat = []
        st.session_state.n_calculos = 0
        st.session_state.pagina = "💬 Asistente"
        st.rerun()
    st.download_button("⬇️ Descargar conversación", data=conversacion_markdown(st.session_state.chat, AJ),
                       file_name=f"conversacion_{datetime.now():%Y%m%d_%H%M}.md", mime="text/markdown",
                       use_container_width=True, disabled=not st.session_state.chat)
    c1, c2 = st.columns(2)
    c1.metric("Consultas", sum(1 for r, *_ in st.session_state.chat if r == "user"))
    c2.metric("Cálculos", st.session_state.n_calculos)
    st.metric("Documentos en la base", len([a for a in BASE.archivos.values() if not a.get("error")]))
    if CLAVE_ADMIN:
        with st.expander("🔐 Administrador" + (" (activo)" if ES_ADMIN else "")):
            if ES_ADMIN:
                if st.button("Cerrar sesión de administrador", use_container_width=True):
                    st.session_state.admin = False
                    st.rerun()
            else:
                clave = st.text_input("Contraseña", type="password", key="clave_admin")
                if st.button("Entrar", use_container_width=True):
                    if clave == CLAVE_ADMIN:
                        st.session_state.admin = True
                        st.rerun()
                    else:
                        st.error("Contraseña incorrecta")
    color = "#22c55e" if CONECTADO else "#ef4444"
    estado_ia = "IA en la nube" if EN_NUBE else ("Ollama conectado" if CONECTADO else "Ollama desconectado")
    privacidad = ("Las preguntas se procesan con Gemini de Google." if EN_NUBE
                  else "Tus datos y documentos no salen de tu computador.")
    st.markdown(f'<div class="cb-pie"><span class="cb-dot" style="background:{color}"></span>'
                f'{estado_ia} · {MODELO}<br>{privacidad}<br>'
                f'Los números los calcula Python; la IA los interpreta.</div>', unsafe_allow_html=True)


def encabezado(subtitulo: str):
    empresa = f" · {AJ['nombre_empresa']}" if AJ["nombre_empresa"] else ""
    color = "#22c55e" if CONECTADO else "#ef4444"
    estado = "En línea" if EN_NUBE else ("Conectado" if CONECTADO else "Sin conexión")
    st.markdown(f"""<div class="cb-header">{img_svg(LOGO_HEADER, 54)}
      <div><h1>{AJ['nombre_asistente']}</h1><div class="cb-sub">{subtitulo}{empresa}</div></div>
      <div class="cb-der"><span class="cb-pill"><span class="cb-dot" style="background:{color}"></span>
      {estado} · {MODELO}</span>{img_logo_empresa(44)}</div></div>""", unsafe_allow_html=True)


# ---------------------------------------------------------------- página: Asistente
def pagina_asistente():
    encabezado("Mi agente de inteligencia artificial · Ingeniería de Confiabilidad")
    if not CONECTADO and not EN_NUBE:
        st.warning("No hay conexión con Ollama. Ábrelo (ícono de la llama junto al reloj) o revisa el servidor en ⚙️ Ajustes.")

    if not st.session_state.chat:
        st.markdown("""<div class="cb-bienvenida"><h3>¿En qué te ayudo hoy?</h3>
            <p>Pega tus datos de falla o elige un ejemplo para empezar</p></div>""", unsafe_allow_html=True)
        cols = st.columns(3)
        for i, (titulo, texto) in enumerate(EJEMPLOS):
            if cols[i % 3].button(titulo, key=f"ej{i}", use_container_width=True, help=texto):
                st.session_state.pendiente = texto
                st.rerun()

    for rol, texto, calculos, fuentes in st.session_state.chat:
        with st.chat_message(rol, avatar="🧑‍🔧" if rol == "user" else "🛠️"):
            mostrar_fuentes(fuentes)
            if calculos and AJ["mostrar_calculos"]:
                mostrar_calculos(calculos)
            st.markdown(texto)


def responder(pregunta: str):
    st.session_state.chat.append(("user", pregunta, [], []))
    with st.chat_message("user", avatar="🧑‍🔧"):
        st.markdown(pregunta)
    with st.chat_message("assistant", avatar="🛠️"):
        agente = st.session_state.agente
        agente.registro = []
        fuentes = []
        with st.spinner("Analizando..."):
            try:
                if AJ["usar_conocimiento"] and BASE.fragmentos:
                    fuentes = BASE.buscar(pregunta, int(AJ["fragmentos"]), AJ["servidor"], MODELO_EMB)
                respuesta = agente.preguntar_con_conocimiento(pregunta, fuentes)
            except urllib.error.HTTPError as e:
                detalle = e.read().decode("utf-8", "ignore")
                if EN_NUBE and e.code in (500, 502, 503, 504):
                    respuesta = ("⏳ El servicio de IA de Google está saturado en este momento. "
                                 "Espera un minuto y vuelve a enviar tu pregunta.")
                elif EN_NUBE and e.code == 429:
                    respuesta = ("⏳ Hay muchas preguntas al mismo tiempo y se alcanzó el límite gratuito por minuto. "
                                 "Espera unos segundos y vuelve a preguntar.")
                elif EN_NUBE:
                    respuesta = f"⚠️ Error del servicio de IA ({e.code}). Avísale al administrador de la página."
                    print("Error Gemini:", e.code, detalle[:500])
                elif "not found" in detalle.lower():
                    respuesta = (f"⚠️ El modelo **{MODELO}** no está descargado. Ejecuta en la terminal: "
                                 f"`ollama pull {MODELO}`")
                else:
                    respuesta = f"⚠️ Error de Ollama {e.code}: {detalle}"
                agente.mensajes.pop()
            except urllib.error.URLError:
                respuesta = ("⚠️ No me pude conectar con el servicio de IA. Intenta de nuevo en un momento."
                             if EN_NUBE else "⚠️ No me pude conectar con Ollama. Verifica que esté abierto "
                             "o revisa la dirección del servidor en ⚙️ Ajustes.")
                agente.mensajes.pop()
        mostrar_fuentes(fuentes)
        if agente.registro and AJ["mostrar_calculos"]:
            mostrar_calculos(agente.registro)
        st.markdown(respuesta)
    st.session_state.n_calculos += len(agente.registro)
    st.session_state.chat.append(("assistant", respuesta, list(agente.registro), fuentes))


# ---------------------------------------------------------------- página: Ajustes
def pagina_ajustes():
    encabezado("Ajustes del asistente")
    if st.session_state.pop("aviso", None):
        st.success("✅ Ajustes guardados. Se aplican desde la próxima pregunta.")

    with st.form("form_ajustes"):
        with st.container(border=True):
            st.subheader("🏢 Identidad")
            c1, c2 = st.columns(2)
            nombre = c1.text_input("Nombre del asistente", AJ["nombre_asistente"])
            empresa = c2.text_input("Nombre de la empresa", AJ["nombre_empresa"], placeholder="Ej.: Terminal Marítimo S.A.")
            c3, c4, c5 = st.columns([1, 1, 2])
            principal = c3.color_picker("Color principal", AJ["color_principal"])
            acento = c4.color_picker("Color de acento", AJ["color_acento"])
            logo = c5.file_uploader("Logo de la empresa (PNG o JPG)", type=["png", "jpg", "jpeg"])

        with st.container(border=True):
            st.subheader("🧠 Modelo de IA")
            c1, c2 = st.columns(2)
            if EN_NUBE:
                modelo = c1.text_input("Modelo de Gemini", AJ["modelo_gemini"],
                                       help="Modelo de Google usado en la nube, p. ej. gemini-3.8-flash.")
            elif MODELOS:
                opciones = MODELOS if AJ["modelo"] in MODELOS else [AJ["modelo"]] + MODELOS
                modelo = c1.selectbox("Modelo instalado en Ollama", opciones, index=opciones.index(AJ["modelo"]),
                                      help="Modelos más grandes (14b, 32b) razonan mejor pero responden más lento.")
            else:
                modelo = c1.text_input("Modelo de Ollama", AJ["modelo"], help="No se pudo leer la lista de modelos de Ollama.")
            servidor = AJ["servidor"] if EN_NUBE else c2.text_input("Servidor de Ollama", AJ["servidor"])
            c1, c2 = st.columns(2)
            temperatura = c1.slider("Creatividad (temperatura)", 0.0, 1.0, float(AJ["temperatura"]), 0.05,
                                    help="Valores bajos dan respuestas más precisas y consistentes; recomendado 0.1–0.3.")
            contextos = [4096, 8192, 16384, 32768]
            contexto = c2.select_slider("Memoria de contexto (tokens)", contextos, value=int(AJ["contexto"]),
                                        help="Más memoria permite usar más documentos por pregunta, pero consume más RAM.")

        with st.container(border=True):
            st.subheader("💬 Respuestas")
            c1, c2 = st.columns(2)
            niveles = ["Breve", "Normal", "Detallado"]
            nivel = c1.radio("Nivel de detalle", niveles, index=niveles.index(AJ["nivel_detalle"]), horizontal=True)
            unidades = ["horas", "días", "ciclos", "kilómetros"]
            unidad = c2.selectbox("Unidad de tiempo por defecto", unidades, index=unidades.index(AJ["unidad_tiempo"]))
            mostrar = st.toggle("Mostrar el detalle de los cálculos en el chat", AJ["mostrar_calculos"])

        with st.container(border=True):
            st.subheader("📚 Base de conocimiento")
            c1, c2, c3 = st.columns(3)
            usar_con = c1.toggle("Consultar mis documentos", AJ["usar_conocimiento"])
            n_frag = c2.slider("Fragmentos por pregunta", 2, 8, int(AJ["fragmentos"]),
                               help="Cuántos trozos de tus documentos se le pasan al agente en cada pregunta.")
            modelo_emb = c3.text_input("Modelo de búsqueda", AJ["modelo_embeddings"],
                                       help="Modelo de Ollama para búsqueda por significado. Instálalo con: ollama pull nomic-embed-text")

        with st.container(border=True):
            st.subheader("📝 Instrucciones de la empresa")
            extra = st.text_area(
                "Instrucciones y reglas propias", AJ["instrucciones_extra"], height=170,
                placeholder="Ej.:\n- Nuestra matriz de criticidad: A = parada de planta, B = pérdida parcial, C = sin impacto.\n"
                            "- Escala de severidad FMEA: 10 = riesgo a personas...\n- Equipos críticos: P-101A, K-201...")

        if st.form_submit_button("💾 Guardar ajustes", use_container_width=True):
            nuevos = {"nombre_asistente": nombre.strip() or AJUSTES_DEFECTO["nombre_asistente"], "nombre_empresa": empresa.strip(),
                      "color_principal": principal, "color_acento": acento,
                      "modelo": AJ["modelo"] if EN_NUBE else modelo.strip(),
                      "modelo_gemini": modelo.strip() if EN_NUBE else AJ["modelo_gemini"],
                      "temperatura": temperatura, "servidor": servidor.strip() or AJUSTES_DEFECTO["servidor"],
                      "nivel_detalle": nivel, "unidad_tiempo": unidad, "mostrar_calculos": mostrar,
                      "instrucciones_extra": extra, "contexto": contexto, "usar_conocimiento": usar_con,
                      "fragmentos": n_frag, "modelo_embeddings": modelo_emb.strip()}
            if logo is not None:
                ARCHIVO_LOGO.write_bytes(logo.getvalue())
            st.session_state.ajustes = nuevos
            guardar_ajustes(nuevos)
            estado_ollama.clear()
            st.session_state.aviso = True
            st.rerun()

    if EN_NUBE:
        st.info("☁️ Versión en la nube: los ajustes se reinician cuando la página se reinicia. Para cambios "
                "permanentes, edita confiabot_ajustes.json en tu repositorio de GitHub.")
    c1, c2, c3 = st.columns(3)
    if c1.button("🔌 Probar conexión", use_container_width=True, disabled=EN_NUBE):
        estado_ollama.clear()
        ok, modelos = estado_ollama(AJ["servidor"])
        if ok:
            st.success(f"Conectado. Modelos instalados: {', '.join(modelos) or 'ninguno'}")
        else:
            st.error("No se pudo conectar con Ollama en " + AJ["servidor"])
    if c2.button("🖼️ Quitar logo de empresa", use_container_width=True, disabled=not ARCHIVO_LOGO.exists()):
        ARCHIVO_LOGO.unlink(missing_ok=True)
        st.rerun()
    if c3.button("↩️ Restablecer valores", use_container_width=True):
        st.session_state.ajustes = dict(AJUSTES_DEFECTO)
        guardar_ajustes(st.session_state.ajustes)
        st.rerun()
    st.caption(f"Los ajustes se guardan en {ARCHIVO_AJUSTES.name}, en la misma carpeta de la página.")


# ---------------------------------------------------------------- página: Acerca de
def pagina_acerca():
    encabezado("Acerca del asistente")
    c1, c2 = st.columns(2)
    with c1, st.container(border=True):
        st.subheader("🧮 Herramientas de cálculo")
        st.markdown("\n".join(f"- {v}" for v in NOMBRES_HERRAMIENTAS.values()))
    with c2, st.container(border=True):
        st.subheader("📐 Métodos y normas")
        st.markdown("- ISO 14224 · Taxonomía y datos de falla\n- RCM · SAE JA1011 / JA1012\n- FMEA · IEC 60812\n"
                    "- Weibull · Rangos de mediana (Bernard) con rangos ajustados de Johnson\n"
                    "- Reemplazo por edad · Modelo de costo mínimo por unidad de tiempo")
    with st.container(border=True):
        st.subheader("🔒 Privacidad")
        if EN_NUBE:
            st.markdown("Esta versión usa Gemini de Google en la nube. No compartas información confidencial "
                        "de tu empresa en las preguntas.")
        else:
            st.markdown("La IA corre en tu propio computador con Ollama: tus datos de falla y tus documentos "
                        "no salen de tu equipo ni se envían a internet.")


# ---------------------------------------------------------------- página: Conocimiento
def pagina_conocimiento():
    encabezado("Base de conocimiento · los documentos con los que aprende tu agente")
    CARPETA_CONOCIMIENTO.mkdir(exist_ok=True)
    validos = {n: a for n, a in BASE.archivos.items() if not a.get("error")}
    c1, c2, c3 = st.columns(3)
    c1.metric("Documentos", len(validos))
    c2.metric("Fragmentos", len(BASE.fragmentos))
    c3.metric("Tipo de búsqueda", "Por significado" if BASE.busqueda_inteligente else "Por palabras clave",
              help="La búsqueda por significado entiende sinónimos y preguntas redactadas distinto. "
                   "Se activa al instalar el modelo de búsqueda: ollama pull " + AJ["modelo_embeddings"])
    if BASE.fragmentos and not BASE.busqueda_inteligente and not EN_NUBE:
        st.info(f"💡 Para una búsqueda más inteligente, ejecuta en la terminal `ollama pull {AJ['modelo_embeddings']}` "
                "y luego toca **🔄 Actualizar índice**.")

    with st.container(border=True):
        st.subheader("➕ Agregar documentos")
        st.caption("PDF, Word (.docx), Excel (.xlsx), CSV, TXT o Markdown: manuales, procedimientos, normas, "
                   "historiales de falla, informes de análisis causa raíz, planes de mantenimiento...")
        with st.form("subir", clear_on_submit=True):
            archivos = st.file_uploader("Arrastra aquí tus archivos", accept_multiple_files=True,
                                        type=[e.strip(".") for e in sorted(EXTENSIONES)])
            enviado = st.form_submit_button("📥 Agregar y procesar", use_container_width=True)
        if enviado and archivos:
            for a in archivos:
                (CARPETA_CONOCIMIENTO / Path(a.name).name).write_bytes(a.getvalue())
            procesar()

    with st.container(border=True):
        st.subheader("📂 Mis documentos")
        if not BASE.archivos:
            st.write("Aún no hay documentos. Súbelos arriba o cópialos directamente en la carpeta:")
        else:
            st.dataframe([{"Documento": n, "Fragmentos": a["fragmentos"], "Estado": a["error"] or "✅ Listo"}
                          for n, a in sorted(BASE.archivos.items())], hide_index=True, use_container_width=True)
            st.caption("También puedes copiar archivos directamente en esta carpeta y tocar 🔄 Actualizar índice:")
        st.code(str(CARPETA_CONOCIMIENTO), language=None)
        c1, c2 = st.columns([2, 1])
        if c1.button("🔄 Actualizar índice", use_container_width=True, type="primary"):
            procesar()
        if BASE.archivos:
            with c2.popover("🗑️ Quitar documento", use_container_width=True):
                quitar = st.selectbox("Documento", sorted(BASE.archivos))
                if st.button("Quitar de la base", type="primary"):
                    (CARPETA_CONOCIMIENTO / quitar).unlink(missing_ok=True)
                    procesar()

    if BASE.fragmentos:
        with st.container(border=True):
            st.subheader("🔎 Probar búsqueda")
            consulta = st.text_input("Escribe una pregunta para ver qué encontraría tu agente",
                                     placeholder="Ej.: ¿cada cuánto se cambia el aceite del compresor?")
            if consulta:
                resultados = BASE.buscar(consulta, int(AJ["fragmentos"]), AJ["servidor"], MODELO_EMB)
                if not resultados:
                    st.write("No se encontró nada relacionado en tus documentos.")
                for i, f in enumerate(resultados, 1):
                    st.markdown(f"**[{i}] {f['doc']}** · {f['loc']} · relevancia {f['puntaje']}")
                    st.caption(f["texto"][:500] + ("…" if len(f["texto"]) > 500 else ""))


def procesar():
    barra = st.progress(0.0, "Revisando documentos...")
    avisos = BASE.actualizar(AJ["servidor"], MODELO_EMB,
                             progreso=lambda v, txt: barra.progress(v, txt))
    barra.empty()
    st.session_state.aviso_conocimiento = avisos
    st.rerun()


# ---------------------------------------------------------------- enrutamiento
pagina = st.session_state.pagina
if "aviso_conocimiento" in st.session_state and pagina == "📚 Conocimiento":
    avisos = st.session_state.pop("aviso_conocimiento")
    if avisos:
        st.warning("Algunos archivos no se pudieron procesar:\n\n" + "\n".join(f"- {a}" for a in avisos))
    else:
        st.toast("✅ Base de conocimiento actualizada")

if pagina == "📚 Conocimiento":
    pagina_conocimiento()
elif pagina == "⚙️ Ajustes":
    pagina_ajustes()
elif pagina == "ℹ️ Acerca de":
    pagina_acerca()
else:
    pagina_asistente()
    pregunta = st.chat_input(f"Pregúntale a {AJ['nombre_asistente']} o pega tus datos de falla...")
    pregunta = pregunta or st.session_state.pendiente
    st.session_state.pendiente = None
    if pregunta:
        responder(pregunta)
